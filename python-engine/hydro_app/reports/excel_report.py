from __future__ import annotations

from pathlib import Path


def write_excel_report(results: dict, path: str | Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Basin Summary"
    summary.append(["Project", results.get("project", {}).get("name", "")])
    summary.append(["Basin Count", results.get("summary", {}).get("basin_count", 0)])
    summary.append(["Total Area (ac)", results.get("summary", {}).get("total_area_acres", 0)])
    summary.append([])
    summary.append(["Basin", "Name", "Condition", "Area (ac)", "Composite CN", "Tc (min)", "Impervious %", "Local Peak (cfs)", "Routed Peak (cfs)", "Downstream", "Status"])

    for basin in results.get("basins", []):
        summary.append([
            basin.get("id", ""),
            basin.get("name", ""),
            basin.get("condition", ""),
            basin.get("area_acres"),
            basin.get("curve_number"),
            basin.get("tc_minutes"),
            basin.get("impervious_percent"),
            basin.get("local_peak_cfs"),
            basin.get("routed_peak_cfs"),
            basin.get("downstream_id", ""),
            basin.get("status", ""),
        ])

    storms = workbook.create_sheet("Storm Summary")
    storms.append(["Active Storm", results.get("active_storm_id", "")])
    storms.append([])
    storms.append(["ID", "Name", "Source", "Location", "ARI (yr)", "Duration", "Depth (in)", "Distribution", "Time Step (min)"])
    for storm in results.get("storms", []):
        storms.append([
            storm.get("id", ""),
            storm.get("name", ""),
            storm.get("source", ""),
            storm.get("location", ""),
            storm.get("ari_years"),
            storm.get("duration", ""),
            storm.get("rainfall_depth_inches"),
            storm.get("distribution", ""),
            storm.get("time_step_minutes"),
        ])

    peak_flows = workbook.create_sheet("Peak Flows")
    peak_flows.append(["Storm", "ARI (yr)", "Depth (in)", "Duration (min)", "Basin", "Local Peak (cfs)", "Routed Peak (cfs)"])
    for peak_flow in results.get("peak_flows", []):
        peak_flows.append([
            peak_flow.get("storm_id", ""),
            peak_flow.get("ari_years"),
            peak_flow.get("rainfall_depth_inches"),
            peak_flow.get("duration_minutes"),
            peak_flow.get("basin_id", ""),
            peak_flow.get("local_peak_cfs"),
            peak_flow.get("routed_peak_cfs"),
        ])

    hydrographs = workbook.create_sheet("Hydrographs")
    hydrographs.append([
        "Hyd No.",
        "Type",
        "Peak Flow (cfs)",
        "Time Interval (min)",
        "Tc (min)",
        "Time to Peak (min)",
        "Volume (cuft)",
        "Inflow Hyd(s)",
        "Maximum Elevation (ft)",
        "Maximum Storage (cuft)",
        "Description",
        "Status",
    ])
    for hydrograph in results.get("hydrographs", []):
        hydrographs.append([
            hydrograph.get("id", ""),
            hydrograph.get("type", ""),
            hydrograph.get("peak_flow_cfs"),
            hydrograph.get("time_interval_minutes"),
            hydrograph.get("tc_minutes"),
            hydrograph.get("time_to_peak_minutes"),
            hydrograph.get("volume_cuft"),
            ", ".join(hydrograph.get("inflow_ids", [])),
            hydrograph.get("maximum_elevation_ft"),
            hydrograph.get("maximum_storage_cuft"),
            hydrograph.get("description", ""),
            hydrograph.get("status", ""),
        ])

    hydrograph_peaks = workbook.create_sheet("Hydrograph Peaks")
    hydrograph_peaks.append(["Storm", "ARI (yr)", "Depth (in)", "Duration (min)", "Hyd No.", "Type", "Peak Flow (cfs)", "Time to Peak (min)", "Volume (cuft)", "Status"])
    for peak_flow in results.get("hydrograph_peak_flows", []):
        hydrograph_peaks.append([
            peak_flow.get("storm_id", ""),
            peak_flow.get("ari_years"),
            peak_flow.get("rainfall_depth_inches"),
            peak_flow.get("duration_minutes"),
            peak_flow.get("hydrograph_id", ""),
            peak_flow.get("hydrograph_type", ""),
            peak_flow.get("peak_flow_cfs"),
            peak_flow.get("time_to_peak_minutes"),
            peak_flow.get("volume_cuft"),
            peak_flow.get("status", ""),
        ])

    validation = workbook.create_sheet("Validation Report")
    validation.append(["Level", "Message"])
    for level in ("errors", "warnings", "info"):
        for message in results.get("validation", {}).get(level, []):
            validation.append([level[:-1], message])

    workbook.save(path)